from pathlib import Path

import fitz
import pytesseract
from PIL import Image
from docx import Document


def extract_pdf_text(file_path: str):
    document = fitz.open(file_path)
    pages = []

    for page_number, page in enumerate(document):
        text = page.get_text("text").strip()

        if not text:
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2))
            image = Image.frombytes(
                "RGB",
                [pixmap.width, pixmap.height],
                pixmap.samples
            )
            text = pytesseract.image_to_string(image).strip()

        pages.append({
            "page": page_number + 1,
            "text": text
        })

    document.close()
    return pages


def extract_docx_text(file_path: str):
    document = Document(file_path)
    paragraphs = []

    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()

        if text:
            paragraphs.append({
                "paragraph": index + 1,
                "text": text
            })

    for table_index, table in enumerate(document.tables):
        for row_index, row in enumerate(table.rows):
            cells = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]

            if cells:
                paragraphs.append({
                    "table": table_index + 1,
                    "row": row_index + 1,
                    "text": " | ".join(cells)
                })

    return paragraphs


def extract_document(file_path: str):
    path = Path(file_path)
    extension = path.suffix.lower()

    if extension == ".pdf":
        content = extract_pdf_text(str(path))
    elif extension == ".docx":
        content = extract_docx_text(str(path))
    else:
        raise ValueError("Only PDF and DOCX files are supported")

    return content